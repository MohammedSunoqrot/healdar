"""
Healdar export -- PDF and Word reports from a RAGAnswer.

Both formats render the answer through formatting.parse_blocks, so headings,
lists, tables and emphasis come out as real document structure instead of
literal Markdown punctuation.

PDF  : reportlab. English uses the built-in Helvetica. Arabic uses the bundled
       IBM Plex Sans Arabic (SIL OFL): arabic_reshaper joins the letters and
       python-bidi puts each line in right-to-left order, since reportlab draws
       left to right and knows neither. (The PDF used to carry the English
       answer for every Arabic question.) If the font cannot be loaded, an
       Arabic answer falls back to its English version rather than failing.
Word : python-docx with full Unicode -- right-to-left where needed.

Sources are passed as (index, source) pairs. The index is the number the
answer cites -- an earlier version renumbered the cited sources from 1, so an
answer citing [2] and [4] got a reference list labelled [1] and [2].
"""

from __future__ import annotations

import html
import io
import logging
import re
from datetime import datetime
from pathlib import Path

import config
import formatting

logger = logging.getLogger(__name__)

ACCENT_HEX = "#3B5BDB"
GREY_HEX = "#6B7489"
RULE_HEX = "#C9D0DD"

_FONT_DIR = Path(__file__).resolve().parent / "assets" / "fonts"
AR_FONT = "HealdarArabic"
AR_FONT_BOLD = "HealdarArabic-Bold"
_ar_ready: bool | None = None


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

def _clean_citations(text: str) -> str:
    """[Source N] -> [N], matching the numbering of the reference list."""
    return config.CITATION_RE.sub(r"[\1]", text)


def _timestamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _file_date() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M")


def _doc_name(filename: str) -> str:
    return Path(filename).stem.replace("_", " ")


def _indexed(sources) -> list[tuple[int, dict]]:
    """Accept [(n, src), ...] or a plain [src, ...] (numbered from 1)."""
    out = []
    for i, item in enumerate(sources or [], start=1):
        if isinstance(item, tuple) and len(item) == 2:
            out.append((int(item[0]), item[1]))
        else:
            out.append((i, item))
    return out


def _excerpt(text: str, limit: int = 260) -> str:
    text = " ".join((text or "").split())
    return text[:limit].rstrip() + ("…" if len(text) > limit else "")


_FOOTER = (
    "For informational purposes only. Always consult the official regulatory "
    "documents and bodies."
)


# ---------------------------------------------------------------------------
# Arabic typesetting
# ---------------------------------------------------------------------------

def arabic_pdf_available() -> bool:
    """Register the bundled Arabic font once. False when it cannot be used."""
    global _ar_ready
    if _ar_ready is None:
        try:
            import arabic_reshaper  # noqa: F401
            import bidi.mirror  # noqa: F401
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont

            pdfmetrics.registerFont(
                TTFont(AR_FONT, str(_FONT_DIR / "IBMPlexSansArabic-Regular.ttf")))
            pdfmetrics.registerFont(
                TTFont(AR_FONT_BOLD, str(_FONT_DIR / "IBMPlexSansArabic-Bold.ttf")))
            _ar_ready = True
        except Exception as exc:  # a missing library, or an LFS pointer instead of the font
            logger.warning("Arabic PDF unavailable, using English instead: %s", exc)
            _ar_ready = False
    return _ar_ready


_LEVEL_RE = re.compile(r"Level\(\s*(\d+),\s*\)")


def _levels(text: str) -> list[int]:
    """Per-character bidi levels. python-bidi reports them per UTF-8 byte."""
    from bidi import get_display

    debug = get_display(text, base_dir="R", debug=True)
    body = debug.split("levels:", 1)[1].split("paragraphs:", 1)[0]
    per_byte = [int(n) for n in _LEVEL_RE.findall(body)]
    out, pos = [], 0
    for ch in text:
        out.append(per_byte[pos])
        pos += len(ch.encode("utf-8"))
    if pos != len(per_byte):
        raise ValueError("bidi level count does not match the text")
    return out


def _to_visual(line: str) -> str:
    """
    One line in display order: the Unicode bidi algorithm's reordering (rule
    L2) and bracket mirroring (rule L4), applied to the levels python-bidi
    computes. Its own get_display() skips mirroring -- "(MDSW)" printed as
    ")MDSW(" -- and its older pure-Python version splits brackets next to
    numbers ("MDR [1]." as "1] MDR]."). This matches Chrome on every case
    tested, nested brackets and "10(3)(a)" included.
    """
    try:
        levels = _levels(line)
    except Exception:  # an unexpected python-bidi debug format: degrade, don't fail
        from bidi.algorithm import get_display
        return get_display(line, base_dir="R")
    from bidi.mirror import MIRRORED

    chars = [MIRRORED.get(c, c) if lv % 2 else c for c, lv in zip(line, levels, strict=True)]
    order = list(range(len(chars)))
    odd = [lv for lv in levels if lv % 2]
    for lvl in range(max(levels, default=0), min(odd, default=1) - 1, -1):
        i = 0
        while i < len(order):
            if levels[order[i]] < lvl:
                i += 1
                continue
            j = i
            while j < len(order) and levels[order[j]] >= lvl:
                j += 1
            order[i:j] = order[i:j][::-1]
            i = j
    return "".join(chars[k] for k in order)


def rtl_lines(text: str, font: str, size: float, width: float) -> list[str]:
    """
    Arabic text as display-ready lines for reportlab.

    Join the letters, wrap to the column in reading order, then reorder each
    line right to left. The order matters: reordering a whole paragraph before
    wrapping puts its last line first.
    """
    import arabic_reshaper
    from reportlab.pdfbase.pdfmetrics import stringWidth

    shaped = arabic_reshaper.reshape(" ".join((text or "").split()))
    lines: list[str] = []
    current = ""
    for word in shaped.split(" "):
        trial = f"{current} {word}" if current else word
        if current and stringWidth(trial, font, size) > width:
            lines.append(current)
            current = word
        else:
            current = trial
    if current:
        lines.append(current)
    return [_to_visual(line) for line in lines]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def to_pdf(
    question: str,
    answer: str,
    sources,
    jurisdiction: str,
    question_original: str = "",
    *,
    question_en: str = "",
    answer_en: str = "",
) -> bytes:
    """
    An Arabic answer is typeset in Arabic, right to left; question_en and
    answer_en are its English version, used as a note and as the fallback when
    the Arabic font cannot be loaded. For an English answer, question_original
    is the question as first asked.
    """
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        HRFlowable,
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    rtl = formatting.is_rtl_text(answer)
    if rtl and not arabic_pdf_available():
        rtl = False
        question_original = question
        question, answer = question_en or question, answer_en or answer

    accent, grey, rule = (colors.HexColor(c) for c in (ACCENT_HEX, GREY_HEX, RULE_HEX))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        rightMargin=2.2 * cm, leftMargin=2.2 * cm,
        topMargin=2 * cm, bottomMargin=2.2 * cm,
        title="Healdar report", author="Healdar",
    )
    base = getSampleStyleSheet()["Normal"]

    def style(name, **kw):
        return ParagraphStyle(name, parent=base, **kw)

    side = TA_RIGHT if rtl else base.alignment
    s_title = style("T", fontSize=22, textColor=accent, alignment=TA_CENTER,
                    fontName="Helvetica-Bold", spaceAfter=4, leading=26)
    s_meta = style("M", fontSize=8, textColor=grey, alignment=TA_CENTER, spaceAfter=10)
    s_h2 = style("H2", fontSize=12, textColor=accent, fontName="Helvetica-Bold",
                 spaceBefore=12, spaceAfter=5)
    s_h3 = style("H3", fontSize=10.5, fontName="Helvetica-Bold", spaceBefore=8, spaceAfter=3)
    s_body = style("B", fontSize=10, leading=15, spaceAfter=6)
    s_cell = style("C", fontSize=8.5, leading=11)
    s_note = style("N", fontSize=9, textColor=grey, leading=13, alignment=side)
    s_ref = style("R", fontSize=9, leading=13, spaceBefore=3, alignment=side)
    s_exc = style("E", fontSize=8, textColor=grey, leftIndent=14, leading=11, alignment=side)
    s_disc = style("D", fontSize=7.5, textColor=grey, alignment=TA_CENTER, spaceBefore=4)
    # Arabic needs a taller line than Latin at the same size.
    s_ar_h2 = style("AH2", fontName=AR_FONT_BOLD, fontSize=13, leading=20, textColor=accent,
                    alignment=TA_RIGHT, spaceBefore=12, spaceAfter=5)
    s_ar_h3 = style("AH3", fontName=AR_FONT_BOLD, fontSize=11, leading=18, alignment=TA_RIGHT,
                    spaceBefore=8, spaceAfter=3)
    s_ar_body = style("AB", fontName=AR_FONT, fontSize=10.5, leading=18, alignment=TA_RIGHT,
                      spaceAfter=6)
    s_ar_item = style("AI", fontName=AR_FONT, fontSize=10.5, leading=18, alignment=TA_RIGHT,
                      rightIndent=14, spaceAfter=3)
    s_ar_cell = style("AC", fontName=AR_FONT, fontSize=9, leading=14, alignment=TA_RIGHT)
    s_ar_exc = style("AE", fontName=AR_FONT, fontSize=8.5, leading=13, textColor=grey,
                     alignment=TA_RIGHT)

    def rl(text: str) -> str:
        return formatting.to_reportlab(text)

    def ar(lines, st: ParagraphStyle, width: float) -> Paragraph:
        """An Arabic paragraph; each logical line is wrapped and reordered for display."""
        visual: list[str] = []
        for line in [lines] if isinstance(lines, str) else lines:
            plain = "".join(seg.text for seg in formatting.inline_segments(line))
            # The page frame keeps 6 pt of padding on each side; a line even a
            # little too wide gets re-wrapped by reportlab, which moves the
            # line's *start* ("1.") onto a line of its own.
            usable = width - st.leftIndent - st.rightIndent - 16
            visual += rtl_lines(plain, st.fontName, st.fontSize, usable)
        return Paragraph("<br/>".join(html.escape(v, quote=False) for v in visual), st)

    def excerpt(text: str):
        if formatting.is_pdf_renderable(text):
            return Paragraph(f"“{rl(_excerpt(text))}”", s_exc)
        if arabic_pdf_available():
            return ar(_excerpt(text), s_ar_exc, doc.width - 14)
        return Paragraph("Source text is in Arabic -- see the Word export for the original.", s_exc)

    story = [
        Paragraph("Healdar", s_title),
        Paragraph(
            f"Health AI Regulatory Intelligence &nbsp;|&nbsp; Generated {_timestamp()} "
            f"&nbsp;|&nbsp; {rl(jurisdiction)} &nbsp;|&nbsp; v{config.APP_VERSION}",
            s_meta,
        ),
        HRFlowable(width="100%", thickness=0.8, color=rule, spaceAfter=8),
    ]

    # ── Question ─────────────────────────────────────────────────────────
    if rtl:
        story += [ar("السؤال", s_ar_h2, doc.width), ar(question, s_ar_body, doc.width)]
        if question_en and question_en.strip() != question.strip():
            story.append(Paragraph(f"(English) {rl(question_en)}", s_note))
    else:
        story.append(Paragraph("Question", s_h2))
        if (question_original and question_original != question
                and formatting.is_pdf_renderable(question_original)):
            story.append(Paragraph(rl(question_original), s_body))
            story.append(Paragraph(f"(English) {rl(question)}", s_note))
        else:
            story.append(Paragraph(rl(question), s_body))
            if question_original and question_original != question:
                story.append(Paragraph("Asked in Arabic; shown in English translation.", s_note))

    # ── Answer ───────────────────────────────────────────────────────────
    story.append(ar("الإجابة", s_ar_h2, doc.width) if rtl else Paragraph("Answer", s_h2))
    for block in formatting.parse_blocks(_clean_citations(answer)):
        if block.kind == "h":
            story.append(ar(block.lines[0], s_ar_h3, doc.width) if rtl
                         else Paragraph(rl(block.lines[0]), s_h3))
        elif block.kind in {"ul", "ol"}:
            if rtl:
                for k, item in enumerate(block.lines):
                    mark = "•" if block.kind == "ul" else f"{block.start + k}."
                    story.append(ar(f"{mark} {item}", s_ar_item, doc.width))
                continue
            items = [ListItem(Paragraph(rl(i), s_body), leftIndent=12) for i in block.lines]
            story.append(ListFlowable(
                items, bulletType="bullet" if block.kind == "ul" else "1",
                leftIndent=14, bulletFontSize=8 if block.kind == "ul" else 10,
                start=block.start if block.kind == "ol" else None,
            ))
        elif block.kind == "table":
            ncols = max(len(r) for r in block.rows)
            width = doc.width / max(1, ncols)
            rows = [row + [""] * (ncols - len(row)) for row in block.rows]
            if rtl:   # first column on the right
                data = [[ar(c, s_ar_cell, width - 12) for c in row[::-1]] for row in rows]
            else:
                data = [[Paragraph(rl(c), s_cell) for c in row] for row in rows]
            tbl = Table(data, colWidths=[width] * ncols, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, rule),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF1F8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story += [tbl, Spacer(1, 6)]
        elif rtl:
            story.append(ar(block.lines, s_ar_body, doc.width))
        else:
            story.append(Paragraph("<br/>".join(rl(ln) for ln in block.lines), s_body))

    # ── References ───────────────────────────────────────────────────────
    indexed = _indexed(sources)
    if indexed:
        story += [Spacer(1, 6), HRFlowable(width="100%", thickness=0.4, color=rule),
                  ar("المصادر", s_ar_h2, doc.width) if rtl else Paragraph("References", s_h2)]
        for n, src in indexed:
            label = (f"<b>[{n}]</b> {rl(_doc_name(src['filename']))} &nbsp;·&nbsp; "
                     f"{rl(str(src.get('jurisdiction', '')))} &nbsp;·&nbsp; p.{src.get('page_number', '')}")
            story.append(Paragraph(label, s_ref))
            if src.get("text"):
                story.append(excerpt(src["text"]))

    story += [
        Spacer(1, 10), HRFlowable(width="100%", thickness=0.4, color=rule),
        Paragraph(_FOOTER, s_disc),
        Paragraph(f"Generated by Healdar v{config.APP_VERSION} ({config.RELEASE_DATE}) "
                  "· Developed by Mohammed R. S. Sunoqrot", s_disc),
    ]
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def to_docx(
    question: str,
    answer: str,              # the answer as displayed (may be Arabic)
    sources,
    jurisdiction: str,
    lang: str = "en",
) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    accent = RGBColor(0x3B, 0x5B, 0xDB)
    grey = RGBColor(0x6B, 0x74, 0x89)
    rtl = lang == "ar"
    align = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    center = WD_ALIGN_PARAGRAPH.CENTER

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.3)
        sec.left_margin = sec.right_margin = Cm(2.6)

    def set_rtl(par):
        if rtl:
            par._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
            par.alignment = align

    def add_rule():
        p = doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "C9D0DD")):
            bottom.set(qn(k), v)
        border.append(bottom)
        ppr.append(border)

    def add_runs(par, line: str):
        for seg in formatting.inline_segments(line):
            run = par.add_run(seg.text)
            run.bold = seg.bold or None
            run.italic = seg.italic or None
            if seg.code:
                run.font.name = "Consolas"
            if rtl:
                run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))

    def heading(text: str, level: int = 2):
        h = doc.add_heading(text, level=level)
        for r in h.runs:
            r.font.color.rgb = accent
        set_rtl(h)

    title = doc.add_heading("Healdar", 0)
    title.alignment = center
    for r in title.runs:
        r.font.color.rgb = accent
    meta = doc.add_paragraph(
        f"Health AI Regulatory Intelligence  |  Generated {_timestamp()}  |  "
        f"{jurisdiction}  |  v{config.APP_VERSION}"
    )
    meta.alignment = center
    meta.runs[0].font.size = Pt(9)
    meta.runs[0].font.color.rgb = grey
    add_rule()

    heading("السؤال" if rtl else "Question")
    q = doc.add_paragraph()
    add_runs(q, question)
    set_rtl(q)

    heading("الإجابة" if rtl else "Answer")
    for block in formatting.parse_blocks(_clean_citations(answer)):
        if block.kind == "h":
            p = doc.add_paragraph()
            add_runs(p, block.lines[0])
            for r in p.runs:
                r.bold = True
            set_rtl(p)
        elif block.kind in {"ul", "ol"}:
            for item in block.lines:
                p = doc.add_paragraph(style="List Bullet" if block.kind == "ul" else "List Number")
                add_runs(p, item)
                set_rtl(p)
        elif block.kind == "table":
            ncols = max(len(r) for r in block.rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = "Table Grid"
            for r_i, row in enumerate(block.rows):
                cells = table.add_row().cells
                for c_i, text in enumerate(row[:ncols]):
                    par = cells[c_i].paragraphs[0]
                    add_runs(par, text)
                    if r_i == 0:
                        for run in par.runs:
                            run.bold = True
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            for i, line in enumerate(block.lines):
                if i:
                    p.add_run().add_break()
                add_runs(p, line)
            set_rtl(p)

    indexed = _indexed(sources)
    if indexed:
        add_rule()
        heading("المصادر" if rtl else "References")
        for n, src in indexed:
            ref = doc.add_paragraph()
            ref.add_run(f"[{n}]  ").bold = True
            ref.add_run(f"{_doc_name(src['filename'])}  ·  {src.get('jurisdiction', '')}  ·  "
                        f"p.{src.get('page_number', '')}")
            if src.get("text"):
                exc = doc.add_paragraph(f"“{_excerpt(src['text'])}”")
                exc.paragraph_format.left_indent = Cm(1.0)
                exc.runs[0].italic = True
                exc.runs[0].font.size = Pt(9)
                exc.runs[0].font.color.rgb = grey

    add_rule()
    disc = doc.add_paragraph(
        f"{_FOOTER}\nGenerated by Healdar v{config.APP_VERSION} ({config.RELEASE_DATE})"
        "  ·  Developed by Mohammed R. S. Sunoqrot"
    )
    disc.alignment = center
    disc.runs[0].font.size = Pt(8)
    disc.runs[0].font.color.rgb = grey

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
